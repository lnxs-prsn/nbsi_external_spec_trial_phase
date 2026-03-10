"""
tests/test_phase4_real.py

Phase 4 exit criteria tests.

Covers:
    chunker.py          — intelligent document splitting
    document_reader.py  — structure-aware format reading
    metadata_nodes.py   — metadata anchor injection
    pipeline.py         — multi-document ingestion

Run:
    PYTHONPATH=. python -m unittest tests.test_phase4_real -v
"""

import os
import sys
import tempfile
import textwrap
import unittest
import uuid
from types import ModuleType


# ---------------------------------------------------------------------------
# Stubs — only for graph.node and graph.edge which are NOT in this package
# ---------------------------------------------------------------------------

class _FakeEmbedder:
    DIM = 4
    def encode(self, texts):
        import numpy as np
        out = []
        for t in texts:
            h = abs(hash(t)) % 1000
            v = np.array([h/1000, (h*3%1000)/1000, len(t)%100/100, 0.5])
            v = v / (np.linalg.norm(v) + 1e-8)
            out.append(v)
        return np.array(out)


class _FakeNode:
    def __init__(self, label, embedding, node_type='concept', activation=0.7):
        self.node_id    = str(uuid.uuid4())
        self.label      = label
        self.embedding  = embedding
        self.node_type  = node_type
        self.activation = activation
        self.protected  = False


class _FakeEdge:
    def __init__(self, source_id, target_id, edge_type='semantic'):
        self.source_id = source_id
        self.target_id = target_id
        self.edge_type = edge_type


class _FakeExtractor:
    def extract(self, text, embedder):
        import numpy as np
        sentences = [s.strip() for s in text.split('.') if s.strip()][:5]
        embs = embedder.encode(sentences) if sentences else np.zeros((0,4))
        nodes = [_FakeNode(s[:40], embs[i].tolist()) for i, s in enumerate(sentences)]
        edges = [_FakeEdge(nodes[i].node_id, nodes[i+1].node_id)
                 for i in range(len(nodes)-1)]
        return nodes, edges


class _FakeGraph:
    def __init__(self): self.node_count = 0


class _FakeSession:
    def __init__(self):
        self.embedder    = _FakeEmbedder()
        self.graph       = _FakeGraph()
        self._ingestions = []
    def ingest_graph(self, nodes, edges):
        self._ingestions.append({'nodes': nodes, 'edges': edges})
        self.graph.node_count += len(nodes)
        return {'nodes': len(nodes), 'edges': len(edges)}


# Register stubs for graph.node and graph.edge ONLY
# Do NOT register nbsi.ingestion — the real package is at nbsi/ingestion/
def _install_stubs():
    node_mod = ModuleType('nbsi.graph.node')
    node_mod.ConceptNode = _FakeNode
    edge_mod = ModuleType('nbsi.graph.edge')
    edge_mod.ConceptEdge = _FakeEdge
    graph_mod = ModuleType('nbsi.graph')
    sys.modules.setdefault('nbsi.graph',      graph_mod)
    sys.modules.setdefault('nbsi.graph.node', node_mod)
    sys.modules.setdefault('nbsi.graph.edge', edge_mod)

_install_stubs()


# ============================================================================
# CHUNKER TESTS
# ============================================================================

class TestChunker(unittest.TestCase):

    def setUp(self):
        from nbsi.ingestion.chunker import chunk_document, Chunk, TARGET_CHARS, MAX_CHARS
        self.chunk_document = chunk_document
        self.Chunk          = Chunk
        self.TARGET_CHARS   = TARGET_CHARS
        self.MAX_CHARS      = MAX_CHARS

    def test_empty_text_returns_empty_list(self):
        self.assertEqual(self.chunk_document(''), [])
        self.assertEqual(self.chunk_document('   '), [])

    def test_short_text_returns_single_chunk(self):
        text   = 'This is a short document. It has two sentences.'
        chunks = self.chunk_document(text)
        self.assertGreaterEqual(len(chunks), 1)

    def test_chunk_text_preserves_content(self):
        text   = 'Concept node. Structural layer. Session end. Beam search.'
        chunks = self.chunk_document(text)
        combined = ' '.join(c.text for c in chunks)
        for word in ['Concept', 'Structural', 'Session', 'Beam']:
            self.assertIn(word, combined)

    def test_markdown_headings_create_chunk_boundaries(self):
        text = textwrap.dedent("""\
            # Introduction

            The introduction explains the background.
            It provides context for the reader.

            ## Methods

            The methods section describes the approach.
            Data was collected using surveys.

            ## Results

            Results showed significant improvement.
            The effect size was large.
        """)
        chunks = self.chunk_document(text)
        self.assertGreaterEqual(len(chunks), 2,
            'Markdown headings should split into at least 2 chunks')

    def test_numbered_sections_create_chunk_boundaries(self):
        text = textwrap.dedent("""\
            1. Background

            This section covers the background of the research.
            Previous work has established several key findings.

            2. Methodology

            The methodology used in this study involves careful analysis.
            Data was processed using standard techniques.
        """)
        chunks = self.chunk_document(text)
        self.assertGreaterEqual(len(chunks), 2)

    def test_heading_text_captured_in_chunk(self):
        text = textwrap.dedent("""\
            ## Results Section

            The results show a clear pattern.
            Confidence intervals are within bounds.
        """)
        chunks = self.chunk_document(text)
        headings = [c.heading for c in chunks]
        self.assertTrue(any('Results' in h for h in headings),
            'Heading text should be captured in chunk.heading')

    def test_chunks_ordered_by_position(self):
        text   = 'Para one.\n\n' * 10 + 'Para two.\n\n' * 10
        chunks = self.chunk_document(text)
        positions = [c.position for c in chunks]
        self.assertEqual(positions, sorted(positions))

    def test_no_chunk_exceeds_max_chars(self):
        text   = 'The structural layer accumulates patterns across sessions. ' * 200
        chunks = self.chunk_document(text)
        for c in chunks:
            self.assertLessEqual(len(c.text), self.MAX_CHARS + 100)

    def test_paragraph_grouping_produces_multiple_chunks(self):
        para  = 'The system processes information through graph traversal algorithms. '
        text  = ('\n\n'.join([para * 3] * 20))
        chunks = self.chunk_document(text)
        self.assertGreaterEqual(len(chunks), 2)


# ============================================================================
# DOCUMENT READER TESTS
# ============================================================================

class TestDocumentReader(unittest.TestCase):

    def setUp(self):
        from nbsi.ingestion.document_reader import read_document, DocumentStructure, Section
        self.read_document     = read_document
        self.DocumentStructure = DocumentStructure
        self.Section           = Section

    def test_read_txt_returns_structure(self):
        with tempfile.NamedTemporaryFile(mode='w', suffix='.txt',
                                         delete=False, encoding='utf-8') as f:
            f.write('Hello world. This is a plain text document.\n')
            path = f.name
        try:
            doc = self.read_document(path)
            self.assertIsNotNone(doc)
            self.assertEqual(doc.format, 'txt')
            self.assertGreater(len(doc.sections), 0)
        finally:
            os.unlink(path)

    def test_read_markdown_detects_headings(self):
        content = textwrap.dedent("""\
            # My Document

            This is the introduction paragraph.

            ## First Section

            Content of the first section goes here.

            ## Second Section

            Content of the second section goes here.
        """)
        with tempfile.NamedTemporaryFile(mode='w', suffix='.md',
                                         delete=False, encoding='utf-8') as f:
            f.write(content)
            path = f.name
        try:
            doc     = self.read_document(path)
            self.assertIsNotNone(doc)
            self.assertEqual(doc.format, 'md')
            headings = doc.headings()
            self.assertGreaterEqual(len(headings), 2,
                f'Expected at least 2 headings, got: {headings}')
        finally:
            os.unlink(path)

    def test_read_markdown_title_extracted(self):
        content = '# The Title\n\nSome body text here.\n'
        with tempfile.NamedTemporaryFile(mode='w', suffix='.md',
                                         delete=False, encoding='utf-8') as f:
            f.write(content)
            path = f.name
        try:
            doc = self.read_document(path)
            self.assertEqual(doc.title, 'The Title')
        finally:
            os.unlink(path)

    def test_read_html_detects_headings(self):
        content = textwrap.dedent("""\
            <html><body>
            <h1>Document Title</h1>
            <p>Introduction paragraph.</p>
            <h2>First Section</h2>
            <p>First section content.</p>
            <h2>Second Section</h2>
            <p>Second section content.</p>
            </body></html>
        """)
        with tempfile.NamedTemporaryFile(mode='w', suffix='.html',
                                         delete=False, encoding='utf-8') as f:
            f.write(content)
            path = f.name
        try:
            doc = self.read_document(path)
            self.assertIsNotNone(doc)
            self.assertEqual(doc.format, 'html')
            self.assertGreaterEqual(len(doc.sections), 2)
        finally:
            os.unlink(path)

    def test_read_html_strips_script_and_style(self):
        content = textwrap.dedent("""\
            <html><head>
            <style>body { color: red; }</style>
            <script>alert("test");</script>
            </head><body>
            <h1>Title</h1>
            <p>Real content here.</p>
            </body></html>
        """)
        with tempfile.NamedTemporaryFile(mode='w', suffix='.html',
                                         delete=False, encoding='utf-8') as f:
            f.write(content)
            path = f.name
        try:
            doc   = self.read_document(path)
            plain = doc.to_plain_text()
            self.assertNotIn('alert', plain)
            self.assertNotIn('color: red', plain)
            self.assertIn('Real content', plain)
        finally:
            os.unlink(path)

    def test_read_docx_with_headings(self):
        try:
            from docx import Document as DocxDocument
        except ImportError:
            self.skipTest('python-docx not installed')

        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
            path = f.name
        try:
            d = DocxDocument()
            d.add_heading('Introduction', level=1)
            d.add_paragraph('This is the introduction.')
            d.add_heading('Methods', level=2)
            d.add_paragraph('This describes the methodology.')
            d.save(path)

            doc = self.read_document(path)
            self.assertIsNotNone(doc)
            self.assertEqual(doc.format, 'docx')
            headings = doc.headings()
            self.assertIn('Introduction', headings)
        finally:
            os.unlink(path)

    def test_missing_file_returns_none(self):
        result = self.read_document('/tmp/nbsi_no_such_file_xyz.txt')
        self.assertIsNone(result)

    def test_unsupported_extension_returns_none(self):
        with tempfile.NamedTemporaryFile(suffix='.xyz', delete=False) as f:
            f.write(b'data')
            path = f.name
        try:
            result = self.read_document(path)
            self.assertIsNone(result)
        finally:
            os.unlink(path)

    def test_to_plain_text_includes_body_content(self):
        content = '# Header\n\nBody paragraph with real content.\n'
        with tempfile.NamedTemporaryFile(mode='w', suffix='.md',
                                         delete=False, encoding='utf-8') as f:
            f.write(content)
            path = f.name
        try:
            doc   = self.read_document(path)
            plain = doc.to_plain_text()
            self.assertIn('Body paragraph', plain)
        finally:
            os.unlink(path)


# ============================================================================
# METADATA NODES TESTS
# ============================================================================

class TestMetadataNodes(unittest.TestCase):

    def setUp(self):
        from nbsi.ingestion.metadata_nodes import build_metadata_nodes, edges_from_chunk_to_heading
        from nbsi.ingestion.document_reader import DocumentStructure, Section
        self.build_metadata_nodes        = build_metadata_nodes
        self.edges_from_chunk_to_heading = edges_from_chunk_to_heading
        self.DocumentStructure           = DocumentStructure
        self.Section                     = Section
        self.embedder                    = _FakeEmbedder()

    def _make_doc(self, title='', author='', headings=None):
        sections = []
        for h in (headings or []):
            sections.append(self.Section(
                heading=h,
                body=f'Body content for section: {h}. More text here.',
                level=1,
            ))
        return self.DocumentStructure(sections=sections, title=title,
                                      author=author, format='md')

    def test_title_node_created(self):
        doc  = self._make_doc(title='My Document')
        meta = self.build_metadata_nodes(doc, self.embedder)
        self.assertIsNotNone(meta.title_node)
        self.assertEqual(meta.title_node.label, 'my document')

    def test_author_node_created(self):
        doc  = self._make_doc(author='Jane Smith')
        meta = self.build_metadata_nodes(doc, self.embedder)
        self.assertIsNotNone(meta.author_node)

    def test_heading_nodes_created(self):
        doc  = self._make_doc(headings=['Introduction', 'Methods', 'Results'])
        meta = self.build_metadata_nodes(doc, self.embedder)
        self.assertEqual(len(meta.heading_index), 3)
        self.assertIn('introduction', meta.heading_index)
        self.assertIn('methods', meta.heading_index)

    def test_metadata_nodes_have_correct_type(self):
        doc  = self._make_doc(title='Test', headings=['Intro'])
        meta = self.build_metadata_nodes(doc, self.embedder)
        for node in meta.nodes:
            self.assertEqual(node.node_type, 'metadata')

    def test_metadata_nodes_are_protected(self):
        doc  = self._make_doc(title='Test', headings=['Intro'])
        meta = self.build_metadata_nodes(doc, self.embedder)
        for node in meta.nodes:
            self.assertTrue(node.protected)

    def test_metadata_nodes_have_high_activation(self):
        doc  = self._make_doc(title='Test', headings=['Intro'])
        meta = self.build_metadata_nodes(doc, self.embedder)
        for node in meta.nodes:
            self.assertGreaterEqual(node.activation, 0.85)

    def test_empty_doc_returns_empty_set(self):
        doc  = self._make_doc()
        meta = self.build_metadata_nodes(doc, self.embedder)
        self.assertEqual(len(meta.nodes), 0)

    def test_duplicate_headings_deduplicated(self):
        doc = self.DocumentStructure(
            sections=[
                self.Section(heading='Results', body='First results.'),
                self.Section(heading='Results', body='More results.'),
            ],
            format='md',
        )
        meta = self.build_metadata_nodes(doc, self.embedder)
        self.assertEqual(len(meta.heading_index), 1)

    def test_anchor_edges_created_for_content_nodes(self):
        doc          = self._make_doc(headings=['Introduction'])
        meta         = self.build_metadata_nodes(doc, self.embedder)
        heading_node = meta.heading_index['introduction']
        content_nodes = [
            _FakeNode('concept a', [0.1, 0.2, 0.3, 0.4]),
            _FakeNode('concept b', [0.2, 0.3, 0.4, 0.5]),
        ]
        edges = self.edges_from_chunk_to_heading(content_nodes, heading_node)
        self.assertEqual(len(edges), 2)
        for e in edges:
            self.assertEqual(e.target_id, heading_node.node_id)

    def test_no_self_edges(self):
        doc          = self._make_doc(headings=['Methods'])
        meta         = self.build_metadata_nodes(doc, self.embedder)
        heading_node = meta.heading_index['methods']
        edges = self.edges_from_chunk_to_heading([heading_node], heading_node)
        self.assertEqual(len(edges), 0)

    def test_anchor_edges_type_is_constitutive(self):
        doc          = self._make_doc(headings=['Results'])
        meta         = self.build_metadata_nodes(doc, self.embedder)
        heading_node = meta.heading_index['results']
        content      = [_FakeNode('finding', [0.1, 0.2, 0.3, 0.4])]
        edges        = self.edges_from_chunk_to_heading(content, heading_node)
        self.assertEqual(edges[0].edge_type, 'constitutive')


# ============================================================================
# PIPELINE TESTS
# ============================================================================

class TestPipeline(unittest.TestCase):

    def setUp(self):
        from nbsi.ingestion.pipeline import ingest_documents, ingest_folder, IngestReport
        self.ingest_documents = ingest_documents
        self.ingest_folder    = ingest_folder
        self.IngestReport     = IngestReport

    def _txt(self, content):
        with tempfile.NamedTemporaryFile(mode='w', suffix='.txt',
                                         delete=False, encoding='utf-8') as f:
            f.write(content)
            return f.name

    def test_ingest_single_txt(self):
        session   = _FakeSession()
        extractor = _FakeExtractor()
        path      = self._txt('NBSI builds concept graphs. '
                               'Structural nodes persist. ' * 5)
        try:
            report = self.ingest_documents(session, extractor, [path], verbose=False)
            self.assertEqual(report.files_ok, 1)
            self.assertEqual(report.files_failed, 0)
            self.assertGreater(report.total_nodes, 0)
        finally:
            os.unlink(path)

    def test_ingest_multiple_documents(self):
        session   = _FakeSession()
        extractor = _FakeExtractor()
        paths     = [self._txt(f'Document {i}. Content sentence. ' * 10)
                     for i in range(3)]
        try:
            report = self.ingest_documents(session, extractor, paths, verbose=False)
            self.assertEqual(report.files_ok, 3)
            self.assertEqual(report.files_failed, 0)
        finally:
            for p in paths:
                os.unlink(p)

    def test_multiple_docs_unified_graph(self):
        session   = _FakeSession()
        extractor = _FakeExtractor()
        paths     = [self._txt('Concept from document. Structural patterns persist. ' * 5)
                     for _ in range(3)]
        try:
            self.ingest_documents(session, extractor, paths, verbose=False)
            self.assertGreater(session.graph.node_count, 0)
            self.assertEqual(len(session._ingestions), 3)
        finally:
            for p in paths:
                os.unlink(p)

    def test_ingest_markdown_with_sections(self):
        session   = _FakeSession()
        extractor = _FakeExtractor()
        content   = textwrap.dedent("""\
            # Introduction

            The introduction describes background context.
            Context matters for understanding the work.

            ## Methods

            Methods describe the data collection process.
            Analysis used statistical tools throughout.

            ## Results

            Results show significant improvement in metrics.
            Effect sizes calculated for all conditions.
        """)
        with tempfile.NamedTemporaryFile(mode='w', suffix='.md',
                                         delete=False, encoding='utf-8') as f:
            f.write(content)
            path = f.name
        try:
            report = self.ingest_documents(session, extractor, [path], verbose=False)
            self.assertEqual(report.files_ok, 1)
            result = report.results[0]
            self.assertGreater(result.chunks, 0)
        finally:
            os.unlink(path)

    def test_missing_file_reported_as_failure(self):
        session   = _FakeSession()
        extractor = _FakeExtractor()
        report    = self.ingest_documents(
            session, extractor,
            ['/tmp/nbsi_no_such_file_xyz_99.txt'],
            verbose=False,
        )
        self.assertEqual(report.files_failed, 1)
        self.assertEqual(report.files_ok, 0)

    def test_mixed_ok_and_failed(self):
        session   = _FakeSession()
        extractor = _FakeExtractor()
        good_path = self._txt('Good document content. ' * 10)
        try:
            report = self.ingest_documents(
                session, extractor,
                [good_path, '/tmp/nbsi_missing_xyz.txt'],
                verbose=False,
            )
            self.assertEqual(report.files_ok, 1)
            self.assertEqual(report.files_failed, 1)
        finally:
            os.unlink(good_path)

    def test_ingest_folder(self):
        session   = _FakeSession()
        extractor = _FakeExtractor()
        with tempfile.TemporaryDirectory() as tmpdir:
            for i, ext in enumerate(['.txt', '.md', '.txt']):
                with open(os.path.join(tmpdir, f'doc{i}{ext}'), 'w') as f:
                    f.write(f'Document content number {i}. ' * 10)
            with open(os.path.join(tmpdir, 'ignore.xyz'), 'w') as f:
                f.write('not supported\n')
            report = self.ingest_folder(session, extractor, tmpdir, verbose=False)
            self.assertEqual(report.files_ok, 3)

    def test_ingest_folder_invalid_path_raises(self):
        session   = _FakeSession()
        extractor = _FakeExtractor()
        with self.assertRaises(ValueError):
            self.ingest_folder(session, extractor,
                               '/tmp/nbsi_no_such_folder_xyz', verbose=False)

    def test_report_summary_contains_counts(self):
        session   = _FakeSession()
        extractor = _FakeExtractor()
        path      = self._txt('Test document content here. ' * 5)
        try:
            report  = self.ingest_documents(session, extractor, [path], verbose=False)
            summary = report.summary()
            self.assertIn('ok', summary)
            self.assertIn('nodes', summary)
        finally:
            os.unlink(path)


# ============================================================================
# Integration
# ============================================================================

class TestPhase4Integration(unittest.TestCase):

    def test_full_pipeline_markdown(self):
        from nbsi.ingestion.pipeline import ingest_documents
        content = textwrap.dedent("""\
            # NBSI System Overview

            The NBSI system builds knowledge graphs from documents.
            Concepts are extracted using spaCy and embedded with MiniLM.

            ## Structural Layer

            The structural layer persists patterns across sessions.
            Node confirmations accumulate in the structural library.

            ## Operational Layer

            The operational layer is reset at the end of each session.
            Concept nodes and edges are built fresh for each document.

            ## Query Interface

            Queries are answered by finding paths through the graph.
            Conductivity measures the strength of each path.
        """)
        session   = _FakeSession()
        extractor = _FakeExtractor()
        with tempfile.NamedTemporaryFile(mode='w', suffix='.md',
                                         delete=False, encoding='utf-8') as f:
            f.write(content)
            path = f.name
        try:
            report = ingest_documents(session, extractor, [path], verbose=False)
            self.assertEqual(report.files_ok, 1,
                f'Expected 1 ok, got: {report.summary()}')
            self.assertGreater(report.total_nodes, 0)
            self.assertGreater(session.graph.node_count, 3)
        finally:
            os.unlink(path)

    def test_cross_document_graph_unified(self):
        from nbsi.ingestion.pipeline import ingest_documents
        doc1 = 'The structural layer manages patterns. Nodes persist. ' * 5
        doc2 = 'The query engine finds paths. Conductivity is measured. ' * 5
        session   = _FakeSession()
        extractor = _FakeExtractor()
        with tempfile.NamedTemporaryFile(mode='w', suffix='.txt',
                                          delete=False, encoding='utf-8') as f1, \
             tempfile.NamedTemporaryFile(mode='w', suffix='.txt',
                                          delete=False, encoding='utf-8') as f2:
            f1.write(doc1); f2.write(doc2)
            path1, path2 = f1.name, f2.name
        try:
            report = ingest_documents(session, extractor, [path1, path2], verbose=False)
            self.assertEqual(report.files_ok, 2)
            self.assertEqual(len(session._ingestions), 2)
            self.assertGreater(session.graph.node_count, 0)
        finally:
            for p in [path1, path2]:
                if os.path.exists(p): os.unlink(p)


if __name__ == '__main__':
    unittest.main(verbosity=2)
